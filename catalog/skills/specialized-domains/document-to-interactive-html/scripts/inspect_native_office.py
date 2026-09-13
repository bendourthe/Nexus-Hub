#!/usr/bin/env python3
"""Verify native PPTX and DOCX output for defects the file's own properties hide."""

from __future__ import annotations

import argparse
import json
import platform
import re
import sys
from pathlib import Path
from typing import Any

# Rule owners: pptx-generation/references/native-motion.md owns chart-title
# visibility; docx-generation owns table grid sizing. This module only MEASURES
# them, per the one-owner-per-rule split in AGENTS.md.

# ECMA-376 21.2.2.8: when c:autoTitleDeleted is absent or val="0" and the chart
# carries no explicit c:title, PowerPoint SYNTHESISES a title at render time from
# the single series name. It is drawn unstyled, so on a dark slide it renders as
# near-black on dark. python-pptx reports has_title=False and COM reports
# HasTitle=False, because the title exists in neither the file nor the object
# model -- only in the render. Setting chart.has_title = False writes val="1".
AUTO_TITLE_DELETED = re.compile(r"<c:autoTitleDeleted[^>]*\bval=\"(?P<val>[^\"]+)\"")
EXPLICIT_TITLE = re.compile(r"<c:title\b")

# Word resolves a layout that declared widths disagree with; allow sub-point rounding.
TABLE_OVERFLOW_TOLERANCE_PT = 0.5

EXIT_PASS = 0
EXIT_FINDINGS = 1
EXIT_UNVERIFIED = 2

TRUTHY = {"1", "true", "on"}


class Unavailable(RuntimeError):
    """A required native capability is absent; never report this as a pass."""


def renders_automatic_title(chart_xml: str, series_count: int) -> bool:
    """True when PowerPoint will draw an automatic title the author never styled."""
    if EXPLICIT_TITLE.search(chart_xml):
        return False
    # An automatic title is only generated for a single-series chart.
    if series_count != 1:
        return False
    match = AUTO_TITLE_DELETED.search(chart_xml)
    if match is None:
        return True  # absent defaults to val="0"
    return match.group("val").strip().lower() not in TRUTHY


def inspect_pptx(source: Path) -> dict[str, Any]:
    """Check every chart for an unstyled automatic title. No Office required."""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise Unavailable(f"python-pptx is required to inspect a presentation: {exc}") from exc

    presentation = Presentation(str(source))
    charts: list[dict[str, Any]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart = shape.chart
            try:
                series_count = sum(len(plot.series) for plot in chart.plots)
            except Exception:  # noqa: BLE001 - malformed plot area
                series_count = 0
            automatic = renders_automatic_title(chart._chartSpace.xml, series_count)
            charts.append(
                {
                    "slide": index,
                    "shape": shape.name,
                    "series": series_count,
                    "declares_title": bool(chart.has_title),
                    "renders_automatic_title": automatic,
                    "severity": "fail" if automatic else "ok",
                }
            )

    failures = sum(1 for c in charts if c["severity"] == "fail")
    return {
        "artifact": str(source),
        "kind": "pptx",
        "status": "fail" if failures else "pass",
        "charts": len(charts),
        "failures": failures,
        "detail": charts,
    }


def _require_com() -> Any:
    if platform.system() != "Windows":
        raise Unavailable(f"native Word inspection requires Windows, not {platform.system()}")
    try:
        import win32com.client
    except ImportError as exc:
        raise Unavailable(f"pywin32 is not installed: {exc}") from exc
    return win32com.client


def inspect_docx(source: Path) -> dict[str, Any]:
    """Measure table geometry as Word resolves it, not as the file declares it.

    The recorded defect was a page-3 table reaching 623.64pt on a 595.44pt page
    while its declared header widths disagreed with the grid. Only Word's own
    layout resolves the value that actually prints.
    """
    client = _require_com()
    try:
        app = client.Dispatch("Word.Application")
    except Exception as exc:
        raise Unavailable(f"Word is not available through COM: {exc}") from exc

    app.Visible = False
    document = None
    tables: list[dict[str, Any]] = []
    try:
        # Office COM resolves against its own working directory and fails
        # ERROR_PATH_NOT_FOUND on a relative path; always hand it an absolute one.
        document = app.Documents.Open(
            str(source.resolve()), ReadOnly=True, AddToRecentFiles=False, Visible=False
        )
        usable_by_section = {
            index: section.PageSetup.PageWidth
            - section.PageSetup.LeftMargin
            - section.PageSetup.RightMargin
            for index, section in enumerate(document.Sections, start=1)
        }
        for position in range(1, document.Tables.Count + 1):
            table = document.Tables(position)
            section = table.Range.Information(2)  # wdActiveEndSectionNumber
            page = table.Range.Information(3)  # wdActiveEndPageNumber
            usable = usable_by_section.get(section, max(usable_by_section.values()))
            widths = [float(cell.Width) for cell in table.Rows(1).Cells]
            total = sum(widths)
            overflow = total - usable
            tables.append(
                {
                    "table": position,
                    "page": page,
                    "columns": len(widths),
                    "total_pt": round(total, 2),
                    "usable_pt": round(usable, 2),
                    "overflow_pt": round(overflow, 2),
                    "severity": "fail" if overflow > TABLE_OVERFLOW_TOLERANCE_PT else "ok",
                }
            )
    finally:
        if document is not None:
            document.Close(0)  # wdDoNotSaveChanges
        app.Quit()

    failures = sum(1 for t in tables if t["severity"] == "fail")
    return {
        "artifact": str(source),
        "kind": "docx",
        "status": "fail" if failures else "pass",
        "tables": len(tables),
        "failures": failures,
        "detail": tables,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pptx", type=Path, help="presentation whose charts are checked")
    parser.add_argument("--docx", type=Path, help="document whose tables are measured")
    parser.add_argument("--json", type=Path, dest="json_out", help="write the evidence record")
    args = parser.parse_args(argv)

    if not args.pptx and not args.docx:
        parser.error("supply --pptx and/or --docx")

    results: list[dict[str, Any]] = []
    unverified: list[dict[str, str]] = []

    for source, runner in ((args.pptx, inspect_pptx), (args.docx, inspect_docx)):
        if source is None:
            continue
        if not source.is_file():
            unverified.append({"artifact": str(source), "reason": "file not found"})
            continue
        try:
            results.append(runner(source))
        except Unavailable as exc:
            unverified.append({"artifact": str(source), "reason": str(exc)})

    record = {
        "results": results,
        "unverified": unverified,
        "overflow_tolerance_pt": TABLE_OVERFLOW_TOLERANCE_PT,
    }
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(json.dumps(record, indent=1) + "\n", encoding="utf-8")

    for result in results:
        print(
            f"{result['kind']}: {result['artifact']} -> {result['status']} "
            f"({result['failures']} failures)"
        )
        for item in result["detail"]:
            if item["severity"] != "fail":
                continue
            if result["kind"] == "pptx":
                print(
                    f"  slide {item['slide']}: chart {item['shape']!r} renders an unstyled "
                    f"automatic title ({item['series']} series, autoTitleDeleted not set)"
                )
            else:
                print(
                    f"  table {item['table']} page {item['page']}: {item['total_pt']}pt "
                    f"exceeds {item['usable_pt']}pt usable by {item['overflow_pt']}pt"
                )

    for entry in unverified:
        print(f"unverified: {entry['artifact']} -- {entry['reason']}")

    if unverified:
        return EXIT_UNVERIFIED
    return EXIT_FINDINGS if any(r["status"] == "fail" for r in results) else EXIT_PASS


if __name__ == "__main__":
    sys.exit(main())
