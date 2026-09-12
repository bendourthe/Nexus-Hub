"""Frozen neutral inputs are reproducible evidence, not paid authoring success."""

import hashlib
import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
CORPUS = ROOT / "tests/fixtures/interactive-handbooks-qualification"
EXTRACTOR = (
    ROOT
    / "catalog/skills/specialized-domains/document-to-interactive-html/scripts/extract_content.py"
)


def test_frozen_inputs_have_exact_bytes_and_no_unlisted_files():
    manifest = json.loads((CORPUS / "frozen-manifest.json").read_text())
    actual = {
        path.relative_to(CORPUS).as_posix(): hashlib.sha256(
            path.read_bytes()
        ).hexdigest()
        for path in CORPUS.rglob("*")
        if path.is_file() and path.name != "frozen-manifest.json"
    }
    assert actual == manifest["files"]


@pytest.fixture(scope="module")
def extractor():
    spec = importlib.util.spec_from_file_location("qualification_extractor", EXTRACTOR)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_report_extraction_preserves_conflicting_totals(extractor):
    model = extractor.build_model(
        [
            str(CORPUS / "report/inputs/pilot.docx"),
            str(CORPUS / "report/inputs/observations.csv"),
        ]
    )
    text = json.dumps(model)
    for fact in ("North", "South", "12", "18", "31", "30", "unavailable"):
        assert fact in text
    assert len(model["sources"]) == 2
    # Native document inspection independently confirms the source discrepancy.
    from docx import Document

    source = Document(CORPUS / "report/inputs/pilot.docx")
    paragraphs = " ".join(p.text for p in source.paragraphs)
    assert "31" in paragraphs and "30" in paragraphs


def test_pptx_extractor_retains_five_source_pages_and_exact_chart(extractor):
    from pptx import Presentation

    source = CORPUS / "presentation/inputs/service-review.pptx"
    deck = Presentation(source)
    assert len(deck.slides) == 5
    chart = next(
        shape.chart
        for slide in deck.slides
        for shape in slide.shapes
        if shape.has_chart
    )
    assert [c.label for c in chart.plots[0].categories] == ["East", "Central", "West"]
    assert list(chart.series[0].values) == [10.0, 15.0, 20.0]
    model = extractor.build_model([str(source)])
    text = json.dumps(model)
    for fact in ("East", "Central", "West", "10", "15", "20", "45", "unavailable"):
        assert fact in text
    assert len(model["sources"]) == 1


def test_native_source_change_refreshes_same_url_and_preserves_anchors(tmp_path):
    shutil.copytree(CORPUS / "repo/inputs", tmp_path, dirs_exist_ok=True)
    output = tmp_path / "docs/handbooks/html/operations.html"
    before = output.read_bytes()
    assert b"Retry limit: 2" in before
    builder = tmp_path / "build.py"
    builder_hash = hashlib.sha256(builder.read_bytes()).hexdigest()
    subprocess.run([sys.executable, str(builder)], check=True, timeout=30)
    updated = output.read_bytes()
    assert b"Retry limit: 3" in updated and updated != before
    assert b'id="overview"' in updated and b'id="workflow"' in updated
    subprocess.run([sys.executable, str(builder)], check=True, timeout=30)
    assert output.read_bytes() == updated
    assert hashlib.sha256(builder.read_bytes()).hexdigest() == builder_hash


@pytest.mark.parametrize("autocrlf", ["true", "false"])
def test_evidence_inputs_keep_exact_bytes_on_git_checkout(tmp_path, autocrlf):
    """Do not normalize evidence hashes to conceal checkout byte differences."""
    checkout = tmp_path / "checkout"
    checkout.mkdir()
    subprocess.run(["git", "init", "-q", str(checkout)], check=True, timeout=30)
    paths = [
        "scripts/installer.ps1",
        "docs/handbooks/_sources/overview/source.md",
        "tests/fixtures/interactive-handbooks-qualification/repo/inputs/config.json",
        "tests/fixtures/interactive-handbooks/subject-fixture.png",
    ]
    for name in [".gitattributes", *paths]:
        target = checkout / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / name, target)
    subprocess.run(
        ["git", "-C", str(checkout), "-c", f"core.autocrlf={autocrlf}", "add", "."],
        check=True,
        capture_output=True,
        timeout=30,
    )
    exported = tmp_path / "exported"
    exported.mkdir()
    subprocess.run(
        [
            "git",
            "-C",
            str(checkout),
            "-c",
            f"core.autocrlf={autocrlf}",
            "checkout-index",
            "--all",
            f"--prefix={exported.as_posix()}/",
        ],
        check=True,
        capture_output=True,
        timeout=30,
    )
    assert all(
        (exported / name).read_bytes() == (ROOT / name).read_bytes() for name in paths
    )
