"""Tests for the user-edit-preservation edit guard (v4.13.1 Phase 3).

Every case runs the real script, mostly as a subprocess so a "new session" is a new
process with a different session id, against a throwaway store.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import io
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPT = REPO_ROOT / "catalog" / "skills" / "workflow" / "user-edit-preservation" / "scripts" / "edit_guard.py"

pptx = pytest.importorskip("pptx")
docx = pytest.importorskip("docx")

# A 1x1 PNG, so the deck has a picture to move without any image file on disk.
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f9d0000000049454e44ae426082"
)


def load_module():
    spec = importlib.util.spec_from_file_location("edit_guard", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture
def store(tmp_path: Path) -> Path:
    return tmp_path / "store"


@pytest.fixture
def work(tmp_path: Path) -> Path:
    folder = tmp_path / "work"
    folder.mkdir()
    return folder


def run(store: Path, *args: str, session: str = "s1", cwd: Path | None = None) -> subprocess.CompletedProcess:
    env = {**os.environ, "NEXUS_EDIT_GUARD_DIR": str(store), "NEXUS_EDIT_GUARD_SESSION": session}
    # Run from the project folder, as an agent does: the store must sit outside it.
    return subprocess.run([sys.executable, str(SCRIPT), *args], env=env, capture_output=True, text=True,
                          cwd=cwd or store.parent / "work", timeout=120, check=False)


def make_text(path: Path) -> Path:
    path.write_text("line one\nline two\n", encoding="utf-8")
    return path


def make_pptx(path: Path) -> Path:
    from pptx.util import Inches

    prs = pptx.Presentation()
    for title in ("Q3 Results", "Revenue", "Outlook"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    prs.slides[1].shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(2), Inches(1), Inches(1))
    prs.save(str(path))
    return path


def make_docx(path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(str(path))
    return path


def user_edits(path: Path) -> None:
    """A content change the user makes: one visible text difference per format."""
    if path.suffix == ".pptx":
        prs = pptx.Presentation(str(path))
        prs.slides[2].shapes.title.text = "Outlook (user edit)"
        prs.save(str(path))
    elif path.suffix == ".docx":
        document = docx.Document(str(path))
        document.add_paragraph("A paragraph the user added.")
        document.save(str(path))
    else:
        with path.open("a", encoding="utf-8") as stream:
            stream.write("line added by the user\n")


MAKERS = {"text": ("notes.md", make_text), "pptx": ("deck.pptx", make_pptx), "docx": ("report.docx", make_docx)}


@pytest.mark.parametrize("kind", sorted(MAKERS))
def test_record_check_change_diff_accept_cycle(kind: str, store: Path, work: Path) -> None:
    name, maker = MAKERS[kind]
    path = maker(work / name)
    assert run(store, "record", str(path), "--from", "write").returncode == 0
    assert run(store, "check", str(path)).returncode == 0
    user_edits(path)
    changed = run(store, "check", str(path))
    assert changed.returncode == 3, changed.stdout
    shown = run(store, "diff", str(path))
    assert shown.returncode == 0
    assert "user" in shown.stdout.lower(), shown.stdout
    assert run(store, "accept", str(path)).returncode == 0
    assert run(store, "check", str(path)).returncode == 0


def test_cross_session_change_is_detected(store: Path, work: Path) -> None:
    path = make_pptx(work / "deck.pptx")
    assert run(store, "record", str(path), session="session-a").returncode == 0
    user_edits(path)
    result = run(store, "check", str(path), session="session-b")
    assert result.returncode == 3


def test_record_cannot_release_a_changed_file(store: Path, work: Path) -> None:
    """Only accept (after diff) may re-baseline a file the user changed; record is refused."""
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    user_edits(path)
    for source in ("write", "command"):
        assert run(store, "record", str(path), "--from", source).returncode == 3, source
    assert run(store, "check", str(path)).returncode == 3


def test_record_after_a_clean_check_and_the_agents_own_write_succeeds(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    assert run(store, "check", str(path)).returncode == 0
    path.write_text("the agent's own revision\n", encoding="utf-8")
    assert run(store, "record", str(path), "--from", "write").returncode == 0
    assert run(store, "check", str(path)).returncode == 0


def test_diff_against_the_agents_copy_needs_no_record_and_cannot_release(store: Path, work: Path) -> None:
    """With no record, the agent compares the user's file with its own generated copy."""
    from pptx.util import Inches

    ours = make_pptx(work / "out.pptx")
    theirs = work / "deck.pptx"
    theirs.write_bytes(ours.read_bytes())
    prs = pptx.Presentation(str(theirs))
    next(s for s in prs.slides[1].shapes if s.shape_type == 13).left = Inches(4)
    prs.slides[1].notes_slide.notes_text_frame.text = "USER NOTE: check the EMEA figure"
    prs.save(str(theirs))
    result = run(store, "diff", str(theirs), "--against", str(ours))
    assert result.returncode == 0
    assert "USER NOTE: check the EMEA figure" in result.stdout
    assert "non-text change on slide 2" in result.stdout, "a moved picture must be named beside a note change"
    assert run(store, "accept", str(theirs)).returncode == 4, "diff --against must not release a block"


def test_a_moved_picture_is_a_non_text_change_naming_the_slide_part(store: Path, work: Path) -> None:
    from pptx.util import Inches

    path = make_pptx(work / "deck.pptx")
    run(store, "record", str(path))
    prs = pptx.Presentation(str(path))
    picture = next(s for s in prs.slides[1].shapes if s.shape_type == 13)
    picture.left = Inches(4)
    prs.save(str(path))
    result = run(store, "check", str(path), session="another")
    assert result.returncode == 3
    assert "non-text change" in result.stdout
    assert "ppt/slides/slide2.xml" in result.stdout


def _rewrite_metadata(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        items = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in items:
            if info.filename == "docProps/core.xml":
                data = data.replace(b"</cp:coreProperties>",
                                    b"<cp:lastModifiedBy>AutoSave</cp:lastModifiedBy></cp:coreProperties>")
            target.writestr(info, data)


@pytest.mark.parametrize("kind", ["pptx", "docx"])
def test_an_office_metadata_only_change_is_unchanged(kind: str, store: Path, work: Path) -> None:
    name, maker = MAKERS[kind]
    path = maker(work / name)
    run(store, "record", str(path))
    before = path.read_bytes()
    _rewrite_metadata(path)
    assert path.read_bytes() != before
    assert run(store, "check", str(path)).returncode == 0


def test_a_file_over_ten_megabytes_is_fingerprinted_without_a_copy(store: Path, work: Path) -> None:
    path = work / "big.bin"
    path.write_bytes(b"x" * (11 * 1024 * 1024))
    recorded = run(store, "record", str(path))
    assert recorded.returncode == 0 and "copy cap" in recorded.stdout
    with path.open("r+b") as stream:
        stream.seek(-1, os.SEEK_END)
        stream.write(b"y")
    assert run(store, "check", str(path)).returncode == 3
    assert "no stored copy" in run(store, "diff", str(path)).stdout


def test_an_office_lock_file_means_open(store: Path, work: Path) -> None:
    path = make_pptx(work / "deck.pptx")
    run(store, "record", str(path))
    (work / "~$deck.pptx").write_bytes(b"lock")
    result = run(store, "check", str(path))
    assert result.returncode == 3 and "open or conflicting copy" in result.stdout


def test_a_onedrive_conflict_copy_means_conflicting(store: Path, work: Path) -> None:
    path = make_pptx(work / "deck.pptx")
    run(store, "record", str(path))
    make_pptx(work / "deck-DESKTOP-4F2A9C.pptx")
    result = run(store, "check", str(path))
    assert result.returncode == 3 and "conflict copy deck-DESKTOP-4F2A9C.pptx" in result.stdout


def test_accept_without_a_diff_is_refused(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    user_edits(path)
    refused = run(store, "accept", str(path))
    assert refused.returncode == 3 and "diff" in refused.stdout
    assert run(store, "check", str(path)).returncode == 3  # still blocked


def test_accept_after_the_file_changed_again_is_refused(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    user_edits(path)
    run(store, "diff", str(path))
    user_edits(path)
    assert run(store, "accept", str(path)).returncode == 3


def test_accept_from_another_session_is_refused(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    user_edits(path)
    run(store, "diff", str(path), session="reviewer")
    assert run(store, "accept", str(path), session="someone-else").returncode == 3


def test_recording_a_read_of_a_changed_file_is_refused(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path), "--from", "write")
    user_edits(path)
    assert run(store, "record", str(path), "--from", "read").returncode == 3
    assert run(store, "check", str(path)).returncode == 3  # the change was not laundered


def test_no_record_missing_file_and_corrupt_archive(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    assert run(store, "check", str(path)).returncode == 4
    run(store, "record", str(path))
    path.unlink()
    assert run(store, "check", str(path)).returncode == 2
    broken = work / "broken.pptx"
    broken.write_bytes(b"not a zip")
    assert run(store, "record", str(broken)).returncode == 2


def test_a_symlink_at_the_path_is_refused(store: Path, work: Path) -> None:
    target = make_text(work / "notes.md")
    link = work / "link.md"
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("symlinks need extra privileges on this host")
    assert run(store, "record", str(link)).returncode == 2


def test_an_online_only_placeholder_cannot_be_verified(store: Path, work: Path, monkeypatch, capsys) -> None:
    guard = load_module()
    path = make_text(work / "notes.md")
    monkeypatch.setenv("NEXUS_EDIT_GUARD_DIR", str(store))
    monkeypatch.chdir(work)
    assert guard.main(["record", str(path)]) == 0
    monkeypatch.setattr(guard, "is_placeholder", lambda info: True)
    assert guard.main(["check", str(path)]) == 5
    assert "placeholder" in capsys.readouterr().out


# --- privacy ------------------------------------------------------------------


def test_secret_looking_paths_keep_no_copy(store: Path, work: Path) -> None:
    path = work / ".env.local"
    path.write_text("TOKEN=abc\n", encoding="utf-8")
    result = run(store, "record", str(path))
    assert result.returncode == 0 and "secret-looking" in result.stdout
    assert list((store / "copies").iterdir()) == []


def test_the_store_is_owner_only(store: Path, work: Path) -> None:
    run(store, "record", str(make_text(work / "notes.md")))
    if os.name == "nt":
        acl = subprocess.run(["icacls", str(store)], capture_output=True, text=True, check=False).stdout
        assert "OWNER RIGHTS" in acl
        assert "BUILTIN\\Users" not in acl and "Everyone" not in acl
    else:
        assert (store.stat().st_mode & 0o777) == 0o700
        record = next((store / "records").iterdir())
        assert (record.stat().st_mode & 0o777) == 0o600


@pytest.mark.parametrize("where", ["git", "onedrive", "workspace"])
def test_the_store_refuses_leaky_locations(where: str, tmp_path: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    if where == "git":
        repo = tmp_path / "repo"
        (repo / ".git").mkdir(parents=True)
        bad, cwd = repo / "store", tmp_path
    elif where == "onedrive":
        bad, cwd = tmp_path / "OneDrive - Contoso" / "store", tmp_path
    else:
        bad, cwd = work / "store", work
    result = run(bad, "record", str(path), cwd=cwd)
    assert result.returncode == 2 and "refusing a store" in result.stderr
    assert not (bad / "records").exists()


def test_purge_and_thirty_day_pruning(store: Path, work: Path) -> None:
    keep = make_text(work / "keep.md")
    old = make_text(work / "old.md")
    run(store, "record", str(keep))
    run(store, "record", str(old))
    guard = load_module()
    record_path = store / "records" / f"{guard.key_for(old)}.json"
    record = json.loads(record_path.read_text(encoding="utf-8"))
    record["at"] = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=31)).isoformat()
    record_path.write_text(json.dumps(record), encoding="utf-8")
    assert run(store, "check", str(old)).returncode == 4  # pruned on this run
    assert run(store, "check", str(keep)).returncode == 0
    assert run(store, "purge", "--path", str(keep)).returncode == 0
    assert run(store, "check", str(keep)).returncode == 4


def test_log_lists_accepts_for_the_summary(store: Path, work: Path) -> None:
    path = make_text(work / "notes.md")
    run(store, "record", str(path))
    user_edits(path)
    run(store, "diff", str(path))
    run(store, "accept", str(path))
    assert "accepted" in run(store, "log").stdout and "notes.md" in run(store, "log").stdout


def test_canonical_paths_fold_case_where_the_os_does(work: Path) -> None:
    guard = load_module()
    path = make_text(work / "Notes.md")
    same = guard.key_for(path) == guard.key_for(Path(str(path).swapcase()))
    assert same is (os.name == "nt" or sys.platform == "darwin")
