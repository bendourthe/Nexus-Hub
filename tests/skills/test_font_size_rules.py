"""Guard the standard font-size rules for slide and document generation skills."""

from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
DOMAINS = ROOT / "catalog/skills/specialized-domains"
CHECKER = DOMAINS / "pptx-generation/scripts/check_font_sizes.py"

SLIDE_LIST = "12, 14, 16, 18, 20, 24, 28, 32, 36, 40, 44, 48, 54, 60, 66, 72, 80, 88, 96"
DOCUMENT_LIST = "10, 11, 12, 14, 16, 18, 20, 22, 24, 26, 28, 36, 48, 72"
SLIDE_SIZES = {float(v) for v in SLIDE_LIST.split(", ")}
DOCUMENT_SIZES = {float(v) for v in DOCUMENT_LIST.split(", ")}

spec = importlib.util.spec_from_file_location("check_font_sizes", CHECKER)
checker = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = checker
spec.loader.exec_module(checker)

# Patterns that set a font size in points in the reference examples.
POINT_PATTERNS = [
    re.compile(r"(?:font\.size\s*=|\bsize\s*=|_size\s*:\s*Pt\s*=)\s*Pt\(([\d.]+)\)"),
    re.compile(r"\bfontSize\s*[:=]\s*([\d.]+)"),
    re.compile(r"\.fontSize\(([\d.]+)\)"),
    re.compile(r"\b(?:title|body)Size\s*:\s*([\d.]+)"),
    re.compile(r"setFont\([^,()]+,\s*([\d.]+)\)"),
    re.compile(r"get\(\"size\",\s*([\d.]+)\)"),
    re.compile(r"\"font_size\":\s*([\d.]+)"),
]
CSS_SIZE = re.compile(r"font-size:\s*([\d.]+)\s*([a-z%]+)")
# docx-js TextRun sizes are half-points; table widths also use `size:`.
DOCX_JS_SIZE = re.compile(r"(?:^\s*|TextRun\(\{[^}]*?)\bsize:\s*([\d.]+)")
DOCX_JS_FILE = "step-4-javascript-docx-generation.md"


def _owner(path: str) -> str:
    return (DOMAINS / path / "SKILL.md").read_text(encoding="utf-8")


def _example_sizes(skill: str) -> list[tuple[str, float, str]]:
    """Return (location, points, unit) for every font size in a skill's bundled files."""
    found = []
    for path in sorted((DOMAINS / skill).rglob("*")):
        if path.suffix not in {".md", ".py", ".js"} or path.name == "SKILL.md":
            continue
        for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            where = f"{path.relative_to(DOMAINS)}:{number}"
            for pattern in POINT_PATTERNS:
                found.extend((where, float(v), "pt") for v in pattern.findall(line))
            if "textpath" not in line:  # VML watermark text scales to its shape
                found.extend((where, float(v), unit) for v, unit in CSS_SIZE.findall(line))
            if path.name == DOCX_JS_FILE and "width" not in line:
                found.extend((where, float(v) / 2, "pt") for v in DOCX_JS_SIZE.findall(line))
    return found


def test_slide_owner_states_list_and_minimum():
    body = _owner("pptx-generation")
    assert "## Font Sizes" in body
    assert SLIDE_LIST in body
    assert "below 12 pt" in body and "explicitly asks" in body
    assert "scripts/check_font_sizes.py" in body


def test_document_owner_states_list_minimum_and_body_size():
    body = _owner("docx-generation")
    assert "## Font Sizes" in body
    assert DOCUMENT_LIST in body
    assert "below 10 pt" in body and "explicitly asks" in body
    assert "body text to 11 or 12" in body
    assert "pdf-document-generation" in body


@pytest.mark.parametrize(
    ("skill", "owners"),
    [
        ("pdf-document-generation", ["docx-generation"]),
        ("theme-tokens", ["pptx-generation", "docx-generation"]),
        ("brand-styling", ["pptx-generation", "docx-generation"]),
    ],
)
def test_referencing_skills_name_owner_without_restating(skill, owners):
    body = _owner(skill)
    handoff = [line for line in body.splitlines() if "font size" in line.lower()]
    assert handoff, f"{skill} has no font-size handoff line"
    for owner in owners:
        assert any(f"`{owner}`" in line for line in handoff), f"{skill} does not name {owner}"
    assert SLIDE_LIST not in body and DOCUMENT_LIST not in body
    assert not re.search(r"\b12, 14, 16\b|\b10, 11, 12\b", body), f"{skill} restates a size list"


@pytest.mark.parametrize(
    ("skill", "allowed"),
    [
        ("pptx-generation", SLIDE_SIZES),
        ("docx-generation", DOCUMENT_SIZES),
        ("pdf-document-generation", DOCUMENT_SIZES),
    ],
)
def test_code_examples_use_standard_sizes(skill, allowed):
    sizes = _example_sizes(skill)
    assert len(sizes) >= 10, f"scanner found too few sizes in {skill}; patterns may be stale"
    bad = [f"{w}: {v:g}{u}" for w, v, u in sizes if u != "pt" or v not in allowed]
    assert not bad, "non-standard font sizes:\n" + "\n".join(bad)


def test_checker_rejects_unsupported_file(tmp_path, capsys):
    target = tmp_path / "notes.txt"
    target.write_text("x", encoding="utf-8")
    assert checker.main([str(target)]) == 2
    assert checker.main([]) == 2


def test_checker_flags_off_list_and_small_slide_sizes(tmp_path, capsys):
    pptx = pytest.importorskip("pptx", reason="python-pptx is not installed in this job")
    from pptx.util import Inches, Pt

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text_frame
    for size in (12, 18, 44):
        frame.add_paragraph().add_run().font.size = Pt(size)
    clean = tmp_path / "clean.pptx"
    deck.save(clean)
    assert checker.main([str(clean)]) == 0

    frame.add_paragraph().add_run().font.size = Pt(13)
    frame.add_paragraph().add_run().font.size = Pt(10)
    dirty = tmp_path / "dirty.pptx"
    deck.save(dirty)
    problems = checker.check(dirty)
    assert len(problems) == 2
    assert any("13 pt is not a standard size" in p for p in problems)
    assert any("10 pt is below the 12 pt minimum" in p for p in problems)
    assert checker.main([str(dirty)]) == 1
    assert len(checker.check(dirty, (10.0, 13.0))) == 0


def test_checker_resolves_used_docx_styles_only(tmp_path):
    docx = pytest.importorskip("docx", reason="python-docx is not installed in this job")
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    document = docx.Document()
    document.add_paragraph().add_run("body").font.size = Pt(11)
    clean = tmp_path / "clean.docx"
    document.save(clean)
    assert checker.check(clean) == []

    used = document.styles.add_style("OddBase", WD_STYLE_TYPE.PARAGRAPH)
    used.font.size = Pt(13)
    child = document.styles.add_style("OddChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = used
    unused = document.styles.add_style("Unused", WD_STYLE_TYPE.PARAGRAPH)
    unused.font.size = Pt(15)
    document.add_paragraph("styled", style="OddChild")
    document.add_paragraph().add_run("half").font.size = Pt(10.5)
    dirty = tmp_path / "dirty.docx"
    document.save(dirty)
    problems = checker.check(dirty)
    assert any("'OddChild'" in p and "13 pt" in p for p in problems)
    assert any("10.5 pt is not a standard size" in p for p in problems)
    assert not any("Unused" in p or "15 pt" in p for p in problems)
