"""Exercise the distributed DOCX table example against Word width contracts."""

import re
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

ROOT = Path(__file__).resolve().parents[2]
REFERENCE = (
    ROOT
    / "catalog/skills/specialized-domains/docx-generation/references/step-2-python-python-docx-fundamentals.md"
)


def table_example():
    blocks = re.findall(
        r"```python\n(.*?)\n```", REFERENCE.read_text(encoding="utf-8"), re.DOTALL
    )
    block = next(block for block in blocks if "def add_data_table(" in block)
    namespace = {}
    exec(compile(block, str(REFERENCE), "exec"), namespace)  # noqa: S102 - test the repository-owned example
    return namespace["add_data_table"]


def test_explicit_widths_match_grid_and_every_cell_after_reopen(tmp_path):
    doc = Document()
    section = doc.sections[-1]
    section.page_width = Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.8)
    table_example()(
        doc,
        ["Local source", "Content"],
        [["review-path.svg", "Four directed links including the return"]],
        [2.0, 4.6],
    )
    path = tmp_path / "table.docx"
    doc.save(path)
    reopened = Document(path)
    table = reopened.tables[0]
    expected = [Inches(2.0).twips, Inches(4.6).twips]
    assert not table.autofit
    assert [column.width.twips for column in table.columns] == expected
    assert all(
        [cell.width.twips for cell in row.cells] == expected for row in table.rows
    )


@pytest.mark.parametrize("widths", [[2.0, 8.0], [2.0], [2.0, 0.0]])
def test_invalid_widths_fail_before_adding_a_table(widths):
    doc = Document()
    with pytest.raises(ValueError):
        table_example()(doc, ["Source", "Content"], [["a", "b"]], widths)
    assert not doc.tables
