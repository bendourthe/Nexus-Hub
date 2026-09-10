"""Guard automatic native shape and chart-series builds in retained exports."""

import importlib.util
import sys
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches

OWNER = Path(__file__).resolve().parents[2] / 'catalog/skills/specialized-domains/pptx-generation/scripts/native_motion.py'
spec = importlib.util.spec_from_file_location('pptx_native_motion', OWNER)
motion = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = motion
spec.loader.exec_module(motion)


def specimen():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    shape.text = 'First process step'
    data = CategoryChartData()
    data.categories = ['North', 'South']
    data.add_series('Counts', [4, 7])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(6), Inches(3), data)
    return deck, slide, shape, chart


def test_native_builds_survive_package_roundtrip():
    deck, slide, shape, chart = specimen()
    motion.add_native_fades(slide, [motion.Fade(shape.shape_id, 0, 450), motion.Fade(chart.shape_id, 140, 700, chart_series=0)])
    payload = BytesIO()
    deck.save(payload)
    restored = Presentation(BytesIO(payload.getvalue())).slides[0]
    timing = restored._element.xpath('./p:timing')[0]
    ids = timing.xpath('.//p:cTn/@id')
    assert len(ids) == len(set(ids))
    assert timing.xpath('.//p:cond[@evt="onBegin"]/p:tn/@val') == ['2']
    assert not timing.xpath('.//p:cond[@evt="onClick"]')
    assert timing.xpath('.//p:cTn[@presetClass="entr"]/p:stCondLst/p:cond/@delay') == ['0', '140']
    assert timing.xpath('.//p:animEffect/p:cBhvr/p:cTn/@dur') == ['450', '700']
    assert timing.xpath('.//a:chart/@seriesIdx') == ['0', '0']
    assert timing.xpath('.//a:bldChart/@animBg') == ['0']
    assert list(restored.shapes[1].chart.series[0].values) == [4, 7]


@pytest.mark.parametrize('invalid', [motion.Fade(999), motion.Fade(2, -1), motion.Fade(2, duration_ms=0), motion.Fade(2, chart_series=0), motion.Fade(3, chart_series=9)])
def test_invalid_build_leaves_slide_untouched(invalid):
    _, slide, shape, _ = specimen()
    before = slide._element.xml
    with pytest.raises(ValueError):
        motion.add_native_fades(slide, [motion.Fade(shape.shape_id), invalid])
    assert slide._element.xml == before


def test_existing_animation_is_never_replaced():
    _, slide, shape, _ = specimen()
    motion.add_native_fades(slide, [motion.Fade(shape.shape_id)])
    before = slide._element.xml
    with pytest.raises(ValueError, match='existing'):
        motion.add_native_fades(slide, [motion.Fade(shape.shape_id)])
    assert slide._element.xml == before


def test_empty_build_is_noop():
    _, slide, _, _ = specimen()
    before = slide._element.xml
    motion.add_native_fades(slide, [])
    assert slide._element.xml == before


def test_connector_only_motion_omits_empty_build_list():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(2), Inches(1))
    motion.add_native_fades(slide, [motion.Fade(connector.shape_id)])
    assert not slide._element.xpath('./p:timing/p:bldLst[not(*)]')
    assert slide._element.xpath('./p:timing//p:animEffect/p:cBhvr/p:tgtEl/p:spTgt/@spid') == [str(connector.shape_id)]
