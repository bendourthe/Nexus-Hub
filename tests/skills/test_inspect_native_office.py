"""Guard native PPTX and DOCX verification of defects file properties hide."""

import importlib.util
import platform
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

REPO = Path(__file__).resolve().parents[2]
OWNER = (
    REPO
    / 'catalog/skills/specialized-domains/document-to-interactive-html/scripts/inspect_native_office.py'
)
spec = importlib.util.spec_from_file_location('inspect_native_office', OWNER)
inspector = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = inspector
spec.loader.exec_module(inspector)

ATTEMPTS = REPO / 'docs/releases/v4/v4.9/development/interactive-handbooks/phase-6-native-attempts'


def _deck(path: Path, *, delete_auto_title: bool, series: int = 1) -> Path:
    """Build a neutral single-slide chart deck.

    python-pptx emits autoTitleDeleted val="0" by default and already reports
    has_title as False, so an author who trusts that property never calls the
    setter that writes val="1". That is the defect this fixture reproduces.
    """
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['first', 'second']
    for index in range(series):
        data.add_series(f'Series {index + 1}', (10, 15))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
    )
    if delete_auto_title:
        frame.chart.has_title = False
    presentation.save(str(path))
    return path


class TestRendersAutomaticTitle:
    """The pure rule, independent of any file or Office install."""

    def test_absent_flag_with_one_series_renders_a_title(self):
        assert inspector.renders_automatic_title('<c:chartSpace/>', 1) is True

    def test_val_zero_renders_a_title(self):
        xml = '<c:chartSpace><c:autoTitleDeleted val="0"/></c:chartSpace>'
        assert inspector.renders_automatic_title(xml, 1) is True

    def test_val_one_suppresses_the_title(self):
        xml = '<c:chartSpace><c:autoTitleDeleted val="1"/></c:chartSpace>'
        assert inspector.renders_automatic_title(xml, 1) is False

    @pytest.mark.parametrize('truthy', ['1', 'true', 'True', 'on'])
    def test_truthy_spellings_all_suppress(self, truthy):
        xml = f'<c:chartSpace><c:autoTitleDeleted val="{truthy}"/></c:chartSpace>'
        assert inspector.renders_automatic_title(xml, 1) is False

    def test_explicit_title_is_not_automatic(self):
        xml = '<c:chartSpace><c:title/><c:autoTitleDeleted val="0"/></c:chartSpace>'
        assert inspector.renders_automatic_title(xml, 1) is False

    def test_multi_series_generates_no_automatic_title(self):
        xml = '<c:chartSpace><c:autoTitleDeleted val="0"/></c:chartSpace>'
        assert inspector.renders_automatic_title(xml, 2) is False


class TestInspectPptx:
    def test_default_python_pptx_chart_is_flagged(self, tmp_path):
        result = inspector.inspect_pptx(_deck(tmp_path / 'bad.pptx', delete_auto_title=False))
        assert result['status'] == 'fail'
        assert result['failures'] == 1
        chart = result['detail'][0]
        assert chart['renders_automatic_title'] is True
        # The property an author would trust says there is no title.
        assert chart['declares_title'] is False

    def test_explicitly_deleted_auto_title_passes(self, tmp_path):
        result = inspector.inspect_pptx(_deck(tmp_path / 'good.pptx', delete_auto_title=True))
        assert result['status'] == 'pass'
        assert result['failures'] == 0

    def test_multi_series_chart_passes_without_the_setter(self, tmp_path):
        deck = _deck(tmp_path / 'multi.pptx', delete_auto_title=False, series=2)
        assert inspector.inspect_pptx(deck)['status'] == 'pass'


class TestRetainedRegression:
    """Bind the check to the artifacts whose verdicts Phase 6 already recorded."""

    @pytest.mark.parametrize(
        'relative,expected',
        [
            ('presentation-final/failed-final.pptx', 'fail'),
            ('presentation-final/first-build.pptx', 'fail'),
            ('presentation-original/first-build.pptx', 'pass'),
            ('presentation-repaired/first-build.pptx', 'pass'),
        ],
    )
    def test_recorded_verdicts_reproduce(self, relative, expected):
        artifact = ATTEMPTS / relative
        if not artifact.is_file():
            pytest.skip(f'retained attempt not present: {relative}')
        assert inspector.inspect_pptx(artifact)['status'] == expected


class TestHonestUnavailability:
    def test_missing_artifact_is_unverified_not_a_pass(self, tmp_path, capsys):
        code = inspector.main(['--pptx', str(tmp_path / 'absent.pptx')])
        assert code == inspector.EXIT_UNVERIFIED
        assert 'unverified' in capsys.readouterr().out

    @pytest.mark.skipif(platform.system() == 'Windows', reason='COM is available here')
    def test_docx_off_windows_reports_unavailable(self, tmp_path):
        target = tmp_path / 'x.docx'
        target.write_bytes(b'')
        with pytest.raises(inspector.Unavailable, match='requires Windows'):
            inspector.inspect_docx(target)

    def test_findings_and_pass_use_distinct_exit_codes(self, tmp_path):
        bad = _deck(tmp_path / 'bad.pptx', delete_auto_title=False)
        good = _deck(tmp_path / 'good.pptx', delete_auto_title=True)
        assert inspector.main(['--pptx', str(bad)]) == inspector.EXIT_FINDINGS
        assert inspector.main(['--pptx', str(good)]) == inspector.EXIT_PASS

    def test_evidence_record_is_written(self, tmp_path):
        good = _deck(tmp_path / 'good.pptx', delete_auto_title=True)
        out = tmp_path / 'evidence' / 'record.json'
        inspector.main(['--pptx', str(good), '--json', str(out)])
        assert out.is_file()
        import json

        record = json.loads(out.read_text(encoding='utf-8'))
        assert record['results'][0]['status'] == 'pass'
        assert record['unverified'] == []
