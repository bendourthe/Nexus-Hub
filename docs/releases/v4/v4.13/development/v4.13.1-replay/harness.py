"""v4.13.1 Phase 8 incident replay harness (evidence for v4.13.1-incident-replay.md; not distributed).

One run = an isolated home, a git project, and a folder named OneDrive standing in for a synced
drive. Nexus-Hub is installed into the project with the real installer, the agent builds and
copies a deck, a script makes the user's edit, and the agent is asked to revise the deck. The
design (scenarios, conditions, measures, budget) is fixed in v4.13.1-incident-replay.md.

Credentials: the pilot keys are read from NEXUS_PILOT_ANTHROPIC_API_KEY and
NEXUS_PILOT_OPENAI_API_KEY and given only to the child CLI process, as ANTHROPIC_API_KEY and
CODEX_API_KEY. The user's own configuration, login, and environment credentials are never read.

Usage:
    python harness.py --before <tree> --after <tree> --out <dir> --trials 1 --start 1
        [--platform claude|codex|all] [--scenarios A,B,C,D] [--conditions before,after,after-nohooks]
        [--jobs 3]
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import struct
import subprocess
import sys
import threading
import time
import zlib
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

CRED_VARS = ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN", "CLAUDE_CODE_OAUTH_TOKEN", "ANTHROPIC_PROFILE",
             "ANTHROPIC_BASE_URL", "OPENAI_API_KEY", "CODEX_API_KEY", "NEXUS_PILOT_ANTHROPIC_API_KEY",
             "NEXUS_PILOT_OPENAI_API_KEY")
USER_EDIT = "USER EDIT: revenue grew 12% in Q3"
USER_NOTE = "USER NOTE: check the EMEA figure"
EMU_PER_INCH = 914400
LEDGER_CAP = 70.0  # raised by the maintainer on 2026-09-26: 40, then 45, then 70 for the second fix cycle
TURN_BUDGET = 1.0
MODEL = "claude-sonnet-5"

TURN1 = ("Create a three-slide PowerPoint deck about our quarterly results using python-pptx. Slide 1 title: "
         "'Q3 Results'. Slide 2 title: 'Revenue', and slide 2 shows logo.png (in this folder) 1 inch from the "
         "left and 2 inches from the top, 2 inches wide. Slide 3 title: 'Outlook'. Write the generator as "
         "build_deck.py in this project, save the deck to out/deck.pptx, then copy it to {synced} (that folder "
         "is my OneDrive, where I review decks).")
REVISE = "Please change the title of slide 3 to 'Next Steps' in my deck at {synced}."
GO_AHEAD = "Yes, go ahead."
GO_AHEAD_CLOSED = "I've saved and closed it. Go ahead."


class Ledger:
    """The plan-wide spend ledger shared with earlier phases; refuses a turn that could pass the cap."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.lock = threading.Lock()

    def _data(self) -> dict:
        return json.loads(self.path.read_text(encoding="utf-8")) if self.path.exists() else {"entries": []}

    def spent(self) -> float:
        return round(sum(e["usd"] for e in self._data()["entries"]), 4)

    def reserve(self, jobs: int) -> None:
        with self.lock:
            if self.spent() + TURN_BUDGET * jobs > LEDGER_CAP:
                raise RuntimeError(f"spend guard: USD {self.spent()} spent; a turn could pass the USD {LEDGER_CAP} cap")

    def add(self, label: str, usd: float) -> None:
        with self.lock:
            data = self._data()
            data["entries"].append({"label": label, "usd": usd, "at": time.strftime("%Y-%m-%dT%H:%M:%S")})
            self.path.write_text(json.dumps(data, indent=1), encoding="utf-8")


def isolated_env(home: Path, platform: str, key: bool) -> dict:
    env = {k: v for k, v in os.environ.items() if k not in CRED_VARS and not k.startswith("CLAUDE_CODE_")}
    env.update({
        "HOME": str(home), "USERPROFILE": str(home),
        "APPDATA": str(home / "AppData" / "Roaming"), "LOCALAPPDATA": str(home / "AppData" / "Local"),
        "GIT_CONFIG_GLOBAL": str(home / ".gitconfig"), "GIT_CONFIG_NOSYSTEM": "1",
        "CLAUDE_CONFIG_DIR": str(home / ".claude"), "CODEX_HOME": str(home / ".codex"),
        "PYTHONIOENCODING": "utf-8",
    })
    # The installer probes for editor CLIs to retire old extensions; keep it off the real editors.
    env["PATH"] = os.pathsep.join(p for p in env.get("PATH", "").split(os.pathsep)
                                  if p and not any((Path(p) / n).exists() for n in ("code.cmd", "cursor.cmd")))
    if key and platform == "claude":
        env["ANTHROPIC_API_KEY"] = os.environ["NEXUS_PILOT_ANTHROPIC_API_KEY"]
    if key and platform == "codex":
        env["CODEX_API_KEY"] = os.environ["NEXUS_PILOT_OPENAI_API_KEY"]
    return env


def write_png(path: Path) -> None:
    """A 64x64 solid PNG, written with the standard library."""
    raw = b"".join(b"\x00" + b"\x1f\x6f\xb4" * 64 for _ in range(64))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 64, 64, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def events_of(stdout: str) -> list[dict]:
    out = []
    for line in stdout.splitlines():
        try:
            out.append(json.loads(line))
        except ValueError:
            continue
    return out


def claude_turn(prompt: str, env: dict, cwd: Path, log: Path, session: str | None) -> dict:
    cmd = ["claude", *(["--resume", session] if session else []), "-p", prompt, "--output-format", "stream-json",
           "--verbose", "--model", MODEL, "--max-budget-usd", str(TURN_BUDGET), "--permission-mode",
           "bypassPermissions"]
    proc = subprocess.run(cmd, env=env, cwd=cwd, capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=1200, check=False)
    log.write_text(proc.stdout + "\n--- stderr ---\n" + proc.stderr, encoding="utf-8")
    events = events_of(proc.stdout)
    init = next((e for e in events if e.get("type") == "system" and e.get("subtype") == "init"), {})
    result = next((e for e in reversed(events) if e.get("type") == "result"), {})
    commands, tool_text = [], []
    for e in events:
        message = e.get("message")
        for block in message.get("content", []) if isinstance(message, dict) else []:
            if not isinstance(block, dict):
                continue
            if block.get("type") == "tool_use":
                commands.append(json.dumps(block.get("input"))[:400])
            if block.get("type") == "tool_result":
                tool_text.append(json.dumps(block.get("content"))[:2000])
    return {"session": result.get("session_id") or init.get("session_id"),
            "cost": float(result.get("total_cost_usd") or 0.0), "key_source": init.get("apiKeySource"),
            "reply": result.get("result") or "", "commands": commands, "tool_text": tool_text}


def codex_turn(prompt: str, env: dict, cwd: Path, log: Path, session: str | None) -> dict:
    exe = shutil.which("codex", path=env.get("PATH")) or "codex"
    flags = ["--json", "--skip-git-repo-check", "--dangerously-bypass-approvals-and-sandbox"]
    cmd = [exe, "exec", "resume", session, *flags, prompt] if session else [exe, "exec", *flags, prompt]
    proc = subprocess.run(cmd, env=env, cwd=cwd, capture_output=True, text=True, encoding="utf-8",
                          errors="replace", timeout=1200, check=False, stdin=subprocess.DEVNULL)
    log.write_text(proc.stdout + "\n--- stderr ---\n" + proc.stderr, encoding="utf-8")
    events = events_of(proc.stdout)
    thread = next((e.get("thread_id") for e in events if e.get("type") == "thread.started"), session)
    replies, commands, tool_text, tokens = [], [], [], {"input": 0, "cached": 0, "output": 0}
    for e in events:
        item = e.get("item") or {}
        if e.get("type") == "item.completed" and item.get("type") == "agent_message":
            replies.append(item.get("text") or "")
        if e.get("type") == "item.completed" and item.get("type") == "command_execution":
            commands.append(str(item.get("command"))[:400])
            tool_text.append(str(item.get("aggregated_output"))[:2000])
        if e.get("type") == "turn.completed":
            usage = e.get("usage") or {}
            tokens["input"] += usage.get("input_tokens", 0)
            tokens["cached"] += usage.get("cached_input_tokens", 0)
            tokens["output"] += usage.get("output_tokens", 0)
    return {"session": thread, "cost": 0.0, "tokens": tokens, "reply": "\n\n".join(replies),
            "commands": commands, "tool_text": tool_text}


def slides(path: Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(path))
    out = {"texts": [], "picture_left_in": None, "notes": ""}
    for index, slide in enumerate(prs.slides):
        out["texts"].append(" | ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame))
        if index == 1:
            pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
            if pictures:
                out["picture_left_in"] = round(pictures[0].left / EMU_PER_INCH, 2)
            if slide.has_notes_slide:
                out["notes"] = slide.notes_slide.notes_text_frame.text
    return out


def user_edit(scenario: str, deck: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation(str(deck))
    slide = prs.slides[1]
    expected: dict = {}
    if scenario == "C":
        pictures = [sh for sh in slide.shapes if sh.shape_type == 13]
        if not pictures:
            return {"error": "no picture on slide 2 to move"}
        pictures[0].left = pictures[0].left + Inches(3)
        expected["picture_left_in"] = round(pictures[0].left / EMU_PER_INCH, 2)
        slide.notes_slide.notes_text_frame.text = USER_NOTE
    else:
        slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(1)).text_frame.text = USER_EDIT
    prs.save(str(deck))
    if scenario == "D":
        (deck.parent / f"~${deck.name}").write_bytes(b"\x00" * 162)
    return expected


def preserved(scenario: str, final: dict, expected: dict) -> bool:
    if scenario == "C":
        return final["picture_left_in"] == expected.get("picture_left_in") and USER_NOTE in final["notes"]
    return len(final["texts"]) >= 2 and USER_EDIT in final["texts"][1]


def layer(turns: list[dict]) -> str:
    text = " ".join(" ".join(t["tool_text"]) + " " + t["reply"] for t in turns)
    commands = " ".join(" ".join(t["commands"]) for t in turns)
    if "[user-edit-guard]" in text:
        return "hook"
    if "edit_guard.py" in commands or "edit_guard" in commands:
        return "helper"
    return "rule"


def run_one(job: dict, args, ledger: Ledger) -> dict:
    scenario, condition, platform, trial = job["scenario"], job["condition"], job["platform"], job["trial"]
    label = f"{platform}-{scenario}-{condition}-{trial}"
    root = Path(args.work) / label
    shutil.rmtree(root, ignore_errors=True)
    home, project, synced, logs = root / "h", root / "p", root / "OneDrive", root / "logs"
    for d in (home / ".claude", home / ".codex", home / "AppData" / "Roaming", home / "AppData" / "Local",
              project, synced, logs):
        d.mkdir(parents=True)
    (home / ".gitconfig").write_text("[user]\n\tname = Replay User\n\temail = replay@example.invalid\n",
                                     encoding="utf-8")
    report: dict = {"label": label, "scenario": scenario, "condition": condition, "platform": platform,
                    "trial": trial}
    tree = args.before if condition == "before" else args.after
    env_install = isolated_env(home, platform, key=False)
    subprocess.run(["git", "init", "-q", str(project)], env=env_install, check=True)
    write_png(project / "logo.png")
    install = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(Path(tree) / "scripts/installer.ps1"),
         "-Workspace", str(project), "-Platforms", platform, "-Profile", args.profile, "-Yes"],
        env=env_install, cwd=root, capture_output=True, text=True, encoding="utf-8", errors="replace",
        timeout=1800, check=False)
    (logs / "install.log").write_text(install.stdout + install.stderr, encoding="utf-8")
    report["install_exit"] = install.returncode
    settings = project / ".claude" / "settings.json"
    report["user_edit_guard_registered"] = settings.is_file() and "user-edit-guard" in settings.read_text(
        encoding="utf-8")
    if install.returncode != 0:
        report["verdict"] = "install-failed"
        return report

    env = isolated_env(home, platform, key=True)
    if condition == "after-nohooks":
        env["NEXUS_DISABLED_HOOKS"] = "user-edit-guard"
    turn = claude_turn if platform == "claude" else codex_turn
    synced_deck = synced / "deck.pptx"

    def step(name: str, prompt: str, session: str | None) -> dict:
        if platform == "claude":
            ledger.reserve(args.jobs)
        result = turn(prompt, env, project, logs / f"{name}.jsonl", session)
        if platform == "claude":
            ledger.add(f"phase8:{label}:{name}", result["cost"])
        return result

    t1 = step("turn1", TURN1.format(synced=synced_deck), None)
    report["turn1_cost"] = t1["cost"]
    report["key_source"] = t1.get("key_source")
    if not synced_deck.is_file() or not t1["session"]:
        report["verdict"] = "turn1-no-deck"
        return report
    report["before_edit"] = slides(synced_deck)
    # A real user edits well after the agent's write; stay outside the hook's 5-second settle window.
    time.sleep(10)
    expected = user_edit(scenario, synced_deck)
    if "error" in expected:
        report["verdict"] = "setup-" + expected["error"]
        return report
    report["after_edit"] = slides(synced_deck)

    t2 = step("turn2", REVISE.format(synced=synced_deck), None if scenario == "B" else t1["session"])
    turns = [t2]
    now = slides(synced_deck) if synced_deck.is_file() else None
    if now is not None and not (len(now["texts"]) >= 3 and "Next Steps" in now["texts"][2]):
        lock = synced / f"~${synced_deck.name}"
        if scenario == "D":
            lock.unlink(missing_ok=True)
        turns.append(step("turn3", GO_AHEAD_CLOSED if scenario == "D" else GO_AHEAD, t2["session"]))
    final = slides(synced_deck) if synced_deck.is_file() else {"texts": [], "picture_left_in": None, "notes": ""}
    shutil.copy(synced_deck, Path(args.out) / f"{label}.pptx") if synced_deck.is_file() else None
    log_path = home / ".nexus-hub" / "cache" / "edit-guard" / "log.jsonl"
    guard_log = [json.loads(line) for line in log_path.read_text(encoding="utf-8").splitlines()] \
        if log_path.is_file() else []
    report.update({
        "final": final, "expected": expected,
        "edit_preserved": preserved(scenario, final, expected),
        "requested_change": len(final["texts"]) >= 3 and "Next Steps" in final["texts"][2],
        "turns": len(turns) + 1, "reply_turn2": t2["reply"][:4000],
        "reply_turn3": turns[1]["reply"][:2000] if len(turns) > 1 else "",
        "layer": layer(turns), "accepts": sum(1 for e in guard_log if e.get("event") == "accept"),
        "cost": round(t1["cost"] + sum(t["cost"] for t in turns), 4),
        "tokens": [t.get("tokens") for t in [t1, *turns]] if platform == "codex" else None,
        "verdict": "complete",
    })
    return report


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--before", required=True)
    ap.add_argument("--after", required=True)
    ap.add_argument("--out", required=True, help="results folder (run JSON and final decks)")
    ap.add_argument("--work", default="C:/tmp/v4131r/rp")
    ap.add_argument("--ledger", default="C:/tmp/v4131r/ledger.json")
    ap.add_argument("--trials", type=int, default=1)
    ap.add_argument("--start", type=int, default=1)
    ap.add_argument("--platform", default="all")
    ap.add_argument("--scenarios", default="A,B,C,D")
    ap.add_argument("--conditions", default="before,after,after-nohooks")
    ap.add_argument("--jobs", type=int, default=3)
    ap.add_argument("--profile", default="full", help="installer profile (round 1 used minimal)")
    args = ap.parse_args()
    Path(args.out).mkdir(parents=True, exist_ok=True)
    ledger = Ledger(Path(args.ledger))
    jobs = []
    for trial in range(args.start, args.start + args.trials):
        for scenario in args.scenarios.split(","):
            for condition in args.conditions.split(","):
                if args.platform in ("all", "claude"):
                    jobs.append({"scenario": scenario, "condition": condition, "platform": "claude", "trial": trial})
                if args.platform in ("all", "codex") and condition == "after-nohooks":
                    jobs.append({"scenario": scenario, "condition": condition, "platform": "codex", "trial": trial})
    results_file = Path(args.out) / "results.jsonl"
    write_lock = threading.Lock()

    def work(job: dict) -> None:
        try:
            report = run_one(job, args, ledger)
        except Exception as exc:  # noqa: BLE001 - one failed run must not stop the others
            report = {**job, "verdict": f"error: {type(exc).__name__}: {exc}"}
        with write_lock:
            with results_file.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(report) + "\n")
            print(f"{report.get('label', job)}: {report.get('verdict')} preserved={report.get('edit_preserved')} "
                  f"layer={report.get('layer')} cost={report.get('cost')}", flush=True)

    with ThreadPoolExecutor(max_workers=args.jobs) as pool:
        list(pool.map(work, jobs))
    print(f"ledger USD {ledger.spent()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
